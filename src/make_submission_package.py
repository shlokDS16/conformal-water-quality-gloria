#!/usr/bin/env python
"""Assemble the manuscript submission package in `submission/` (brief 13 section C).

Produces, from files that already exist, without changing any of them:

    submission/
        Manuscript_review.pdf            copy of paper/main.pdf (the review PDF)
        Manuscript.docx                  copy of build/manuscript.docx (the Word manuscript)
        Highlights.docx                  the five highlights as an editable Word file
        Highlights.txt                   the same five lines as plain text
        figures/Figure_1.pdf ... .png    every figure renamed to its PRINTED number, read
        figures/Figure_A1.pdf ...          from paper/main.aux, not from the file name
        latex_source/                    a flat copy of the LaTeX source that compiles alone
        SUBMISSION_CONTENTS.md           what each file is and how it was produced

The figure file names in `figures/` no longer match the printed numbers (figA1_targets.pdf
prints as Figure 2, figA3_wbcoverage_ecdf.pdf as Figure 10), so the mapping is taken from the
label of each `figure` environment and the number `main.aux` gives that label.

`latex_source/` is flat: `\\input{../tables/x.tex}` becomes `\\input{x.tex}` and
`\\includegraphics{../figures/y.pdf}` becomes `\\includegraphics{y.pdf}`, with the table
sources and the figure PDFs copied next to `main.tex`. `--test-compile` compiles that folder
in a scratch directory and checks that the page count matches `paper/main.pdf`.

Usage:
    python src/make_submission_package.py [--out submission] [--test-compile]
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"
FIGDIR = ROOT / "figures"


# --------------------------------------------------------------------- helpers
def balanced(text: str, open_idx: int) -> str:
    depth, i = 0, open_idx
    while i < len(text):
        if text[i] == "\\":
            i += 2
            continue
        if text[i] == "{":
            depth += 1
        elif text[i] == "}":
            depth -= 1
            if depth == 0:
                return text[open_idx + 1:i]
        i += 1
    return text[open_idx + 1:]


def strip_comments(text: str) -> str:
    out, i, n = list(text), 0, len(text)
    while i < n:
        if text[i] == "\\":
            i += 2
            continue
        if text[i] == "%":
            j = i
            while j < n and text[j] != "\n":
                out[j] = " "
                j += 1
            i = j
        else:
            i += 1
    return "".join(out)


def parse_aux_numbers(aux: Path) -> dict[str, str]:
    txt = aux.read_text(encoding="utf-8", errors="replace")
    out: dict[str, str] = {}
    for m in re.finditer(r"\\newlabel\{", txt):
        lab = balanced(txt, m.end() - 1)
        rest = m.end() - 1 + len(lab) + 2
        if rest >= len(txt) or txt[rest] != "{":
            continue
        group = balanced(txt, rest)
        if group.startswith("{"):
            out[lab] = balanced(group, 0)
    return out


def figure_map(tex: Path, aux: Path) -> list[tuple[str, str, Path]]:
    """[(printed number, label, figure file)] in document order."""
    src = strip_comments(tex.read_text(encoding="utf-8", errors="replace"))
    numbers = parse_aux_numbers(aux)
    rows = []
    for m in re.finditer(r"\\begin\{figure\*?\}", src):
        end = src.find("\\end{figure", m.end())
        body = src[m.start():end]
        lab = re.search(r"\\label\{([^}]*)\}", body)
        inc = re.search(r"\\includegraphics(?:\[[^\]]*\])?\{([^}]*)\}", body)
        if not (lab and inc):
            continue
        num = numbers.get(lab.group(1), "?")
        path = (tex.parent / inc.group(1)).resolve()
        rows.append((num, lab.group(1), path))
    return rows


def highlights_lines(tex: Path) -> list[str]:
    src = strip_comments(tex.read_text(encoding="utf-8", errors="replace"))
    m = re.search(r"\\begin\{highlights\}(.*?)\\end\{highlights\}", src, re.S)
    if not m:
        return []
    items = re.split(r"\\item\b", m.group(1))
    return [re.sub(r"\s+", " ", x).strip() for x in items if x.strip()]


def write_highlights_docx(lines: list[str], path: Path) -> None:
    import docx
    from docx.shared import Pt
    d = docx.Document()
    for s in d.sections:
        s.left_margin = s.right_margin = docx.shared.Cm(2.5)
        s.top_margin = s.bottom_margin = docx.shared.Cm(2.5)
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    h = d.add_paragraph()
    r = h.add_run("Highlights")
    r.bold = True
    r.font.size = Pt(14)
    for line in lines:
        p = d.add_paragraph(line, style="List Bullet")
        for run in p.runs:
            run.font.name = "Times New Roman"
            run.font.size = Pt(12)
    # the reviewer must not receive the builder's name in the file
    cp = d.core_properties
    cp.author = ""
    cp.last_modified_by = ""
    cp.title = ""
    cp.category = ""
    cp.comments = ""
    d.save(str(path))


def build_flat_source(out: Path, tex: Path) -> list[str]:
    """Copy main.tex, refs.bib, main.bbl, the table sources and the figures, flattened."""
    out.mkdir(parents=True, exist_ok=True)
    raw = tex.read_text(encoding="utf-8", errors="replace")
    copied: list[str] = []

    def flatten_input(m):
        p = Path(m.group(1))
        src = (tex.parent / p).resolve()
        if src.exists():
            shutil.copy2(src, out / src.name)
            copied.append(src.name)
        return "\\input{%s}" % src.name

    def flatten_graphic(m):
        opt = m.group(1) or ""
        p = Path(m.group(2))
        src = (tex.parent / p).resolve()
        if src.exists():
            shutil.copy2(src, out / src.name)
            copied.append(src.name)
        return "\\includegraphics%s{%s}" % (opt, src.name)

    new = re.sub(r"\\input\{([^}]*)\}", flatten_input, raw)
    new = re.sub(r"\\includegraphics(\[[^\]]*\])?\{([^}]*)\}", flatten_graphic, new)
    with (out / "main.tex").open("w", encoding="utf-8", newline="") as fh:
        fh.write(new)
    for name in ("refs.bib", "main.bbl"):
        s = tex.parent / name
        if s.exists():
            shutil.copy2(s, out / name)
            copied.append(name)
    return sorted(set(copied))


def page_count(pdf: Path) -> int:
    import fitz
    d = fitz.open(str(pdf))
    n = d.page_count
    d.close()
    return n


def test_compile(srcdir: Path) -> tuple[int | None, str]:
    work = Path(tempfile.mkdtemp(prefix="subpkg_")) / "src"
    shutil.copytree(srcdir, work)
    for i in range(3):
        subprocess.run(["pdflatex", "-interaction=nonstopmode", "main.tex"],
                       cwd=str(work), capture_output=True, text=True,
                       encoding="utf-8", errors="replace")
        if i == 0:
            subprocess.run(["bibtex", "main"], cwd=str(work), capture_output=True,
                           text=True, encoding="utf-8", errors="replace")
    pdf = work / "main.pdf"
    log = (work / "main.log").read_text(encoding="utf-8", errors="replace")
    errs = len(re.findall(r"^! ", log, re.M))
    over = len(re.findall(r"^Overfull \\[hv]box", log, re.M))
    note = f"errors {errs}, overfull boxes {over}, directory {work}"
    return (page_count(pdf) if pdf.exists() else None), note


# ------------------------------------------------------------------------ main
def main() -> int:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default=str(ROOT / "submission"))
    ap.add_argument("--tex", default=str(PAPER / "main.tex"))
    ap.add_argument("--pdf", default=str(PAPER / "main.pdf"))
    ap.add_argument("--docx", default=str(ROOT / "build" / "manuscript.docx"))
    ap.add_argument("--test-compile", action="store_true")
    args = ap.parse_args()

    out = Path(args.out)
    tex, pdf, docx_path = Path(args.tex), Path(args.pdf), Path(args.docx)
    out.mkdir(parents=True, exist_ok=True)

    shutil.copy2(pdf, out / "Manuscript_review.pdf")
    shutil.copy2(docx_path, out / "Manuscript.docx")
    print(f"Manuscript_review.pdf  {page_count(pdf)} pages")
    print("Manuscript.docx        copied")

    hl = highlights_lines(tex)
    write_highlights_docx(hl, out / "Highlights.docx")
    with (out / "Highlights.txt").open("w", encoding="utf-8", newline="\r\n") as fh:
        fh.write("Highlights\n\n" + "\n".join(hl) + "\n")
    print(f"Highlights.docx        {len(hl)} bullets")

    figdir = out / "figures"
    if figdir.exists():
        shutil.rmtree(figdir)
    figdir.mkdir(parents=True)
    rows = figure_map(tex, PAPER / "main.aux")
    table = []
    for num, lab, src in rows:
        name = "Figure_" + num.replace(".", "")
        for suffix in (".pdf", ".png"):
            s = src.with_suffix(suffix)
            if s.exists():
                shutil.copy2(s, figdir / (name + suffix))
        table.append((name, num, lab, src.name))
        print(f"  {name:12s} <- {src.name}   (prints as {num})")

    latex_out = out / "latex_source"
    if latex_out.exists():
        shutil.rmtree(latex_out)
    copied = build_flat_source(latex_out, tex)
    print(f"latex_source/          {len(copied) + 1} files, flat paths")

    pages_expected = page_count(pdf)
    pages_got, note = (None, "not run")
    if args.test_compile:
        pages_got, note = test_compile(latex_out)
        ok = pages_got == pages_expected
        print(f"test compile           {pages_got} pages against {pages_expected}"
              f" -> {'OK' if ok else 'MISMATCH'} ({note})")

    lines = [
        "# submission/ contents",
        "",
        "Built by `python src/make_submission_package.py --test-compile`"
        " (brief 13 section C). Every file is a copy of something already built; nothing here"
        " is edited by hand.",
        "",
        "| File | What it is |",
        "|---|---|",
        f"| `Manuscript_review.pdf` | the review PDF, {pages_expected} pages,"
        " `paper/main.pdf` |",
        "| `Manuscript.docx` | the Word manuscript, `build/manuscript.docx`."
        " Do NOT open and re-save it in Word before sending: Word writes the signed-in user's"
        " name into the metadata. If it needs a change, change `paper/main.tex` and run"
        " `python src/make_docx.py` again. |",
        "| `Highlights.docx`, `Highlights.txt` | the five highlights as editable files |",
        "| `figures/` | one file per figure, named by its PRINTED number, vector PDF and"
        " raster PNG |",
        "| `latex_source/` | the LaTeX source with flat paths; compiles on its own |",
        "",
        "## Figure file names",
        "",
        "The names in the project's `figures/` folder no longer match the printed numbers, so"
        " the mapping below is taken from `paper/main.aux`.",
        "",
        "| Submission file | Prints as | Label | Source file |",
        "|---|---|---|---|",
    ]
    for name, num, lab, srcname in table:
        kind = "Fig." if not re.match(r"^[A-Z]\.", num) else "Fig."
        lines.append(f"| `{name}.pdf`, `{name}.png` | {kind} {num} | `{lab}` | `{srcname}` |")
    lines += [
        "",
        "## latex_source/",
        "",
        "`\\input{../tables/x.tex}` became `\\input{x.tex}` and"
        " `\\includegraphics{../figures/y.pdf}` became `\\includegraphics{y.pdf}`; the table"
        " sources, the figure PDFs, `refs.bib` and `main.bbl` sit next to `main.tex`.",
        "",
        f"Test compile: {pages_got} pages against {pages_expected} in the review PDF"
        f" ({note}).",
        "",
        "Files:",
        "",
    ] + [f"- `{n}`" for n in ["main.tex"] + copied]
    (out / "SUBMISSION_CONTENTS.md").write_text("\n".join(lines) + "\n",
                                                encoding="utf-8", newline="")
    print(f"wrote {out / 'SUBMISSION_CONTENTS.md'}")
    if args.test_compile and pages_got != pages_expected:
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
